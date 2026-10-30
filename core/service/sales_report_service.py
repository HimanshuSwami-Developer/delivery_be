from datetime import date, datetime, time, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from django.db.models import F, Sum
from django.db.models.functions import Coalesce

from orders.models import Order, OrderItem
from reports.models import StoreSettings

# The store operates in India — "a day's sales" means an IST calendar day,
# not a UTC one (`settings.TIME_ZONE` is UTC).
STORE_TZ = ZoneInfo("Asia/Kolkata")

PERIODS = ("daily", "weekly", "monthly", "yearly")

# How far back each period looks when no explicit range is given.
_DEFAULT_SPAN = {"daily": 30, "weekly": 12, "monthly": 12, "yearly": 5}

_MONEY_FIELDS = ("subtotal", "discount", "cgst", "sgst", "delivery_fee", "revenue")


def _bucket_start(d: date, period: str) -> date:
    if period == "daily":
        return d
    if period == "weekly":
        return d - timedelta(days=d.weekday())  # Monday
    if period == "monthly":
        return d.replace(day=1)
    return d.replace(month=1, day=1)


def _next_bucket(d: date, period: str) -> date:
    if period == "daily":
        return d + timedelta(days=1)
    if period == "weekly":
        return d + timedelta(days=7)
    if period == "monthly":
        return (d.replace(day=28) + timedelta(days=4)).replace(day=1)
    return d.replace(year=d.year + 1)


def _bucket_label(start: date, period: str) -> str:
    if period == "daily":
        return start.strftime("%d %b %Y")
    if period == "weekly":
        end = start + timedelta(days=6)
        return f"{start.strftime('%d %b')} – {end.strftime('%d %b %Y')}"
    if period == "monthly":
        return start.strftime("%b %Y")
    return str(start.year)


def default_range(period: str, today: date) -> tuple[date, date]:
    """Trailing window ending today: last 30 days / 12 weeks / 12 months /
    5 years, each starting on a bucket boundary."""
    span = _DEFAULT_SPAN[period]
    start = _bucket_start(today, period)
    for _ in range(span - 1):
        start = _bucket_start(start - timedelta(days=1), period)
    return start, today


class SalesReportService:
    """Successfully delivered sales, grouped by day/week/month/year — the
    admin console's "Delivered sales" report and its Excel export. A sale
    is counted on the IST date it was *delivered* (falling back to the
    order's last update for legacy orders without `delivered_at`);
    cancelled and still-in-transit orders are excluded."""

    @staticmethod
    def today() -> date:
        return datetime.now(STORE_TZ).date()

    @staticmethod
    def _delivered_orders(start: date, end: date):
        start_dt = datetime.combine(start, time.min, tzinfo=STORE_TZ)
        end_dt = datetime.combine(end + timedelta(days=1), time.min, tzinfo=STORE_TZ)
        return (
            Order.objects.filter(status=Order.Status.DELIVERED)
            .annotate(sold_at=Coalesce("delivered_at", "updated_at"))
            .filter(sold_at__gte=start_dt, sold_at__lt=end_dt)
        )

    @classmethod
    def build(cls, period: str, start: date, end: date) -> dict:
        orders = list(
            cls._delivered_orders(start, end)
            .select_related("customer")
            .annotate(units=Sum("items__qty"))
            .order_by("sold_at")
        )

        # Every bucket in range, including empty ones, so charts/tables
        # show gaps as zero instead of silently skipping them.
        buckets = {}
        cursor = _bucket_start(start, period)
        while cursor <= end:
            buckets[cursor] = {
                "period_start": cursor.isoformat(),
                "label": _bucket_label(cursor, period),
                "orders": 0,
                "items": 0,
                **{f: 0 for f in _MONEY_FIELDS},
                "cod_revenue": 0,
                "online_revenue": 0,
            }
            cursor = _next_bucket(cursor, period)

        order_rows = []
        for o in orders:
            sold_date = o.sold_at.astimezone(STORE_TZ).date()
            row = buckets[_bucket_start(sold_date, period)]
            row["orders"] += 1
            row["items"] += o.units or 0
            row["subtotal"] += o.subtotal
            row["discount"] += o.discount
            row["cgst"] += o.cgst
            row["sgst"] += o.sgst
            row["delivery_fee"] += o.delivery_fee
            row["revenue"] += o.total
            if o.payment_mode == "cod":
                row["cod_revenue"] += o.total
            else:
                row["online_revenue"] += o.total
            order_rows.append({
                "order_number": o.order_number,
                "delivered_at": o.sold_at.astimezone(STORE_TZ),
                "customer": getattr(o.customer, "name", None) or o.customer.mobile_number,
                "mobile": o.customer.mobile_number,
                "payment_mode": o.get_payment_mode_display(),
                "items": o.units or 0,
                "subtotal": o.subtotal,
                "discount": o.discount,
                "cgst": o.cgst,
                "sgst": o.sgst,
                "delivery_fee": o.delivery_fee,
                "total": o.total,
            })

        rows = list(buckets.values())
        totals = {
            "orders": sum(r["orders"] for r in rows),
            "items": sum(r["items"] for r in rows),
            **{f: sum(r[f] for r in rows) for f in _MONEY_FIELDS},
            "cod_revenue": sum(r["cod_revenue"] for r in rows),
            "online_revenue": sum(r["online_revenue"] for r in rows),
        }
        totals["avg_order_value"] = round(totals["revenue"] / totals["orders"]) if totals["orders"] else 0

        best = max(rows, key=lambda r: r["revenue"], default=None)
        return {
            "period": period,
            "start": start.isoformat(),
            "end": end.isoformat(),
            "totals": totals,
            "best_period": best["label"] if best and best["revenue"] else None,
            "rows": rows,
            "_orders": order_rows,  # Excel only — stripped from the JSON response.
        }

    @classmethod
    def top_products(cls, start: date, end: date):
        return (
            OrderItem.objects.filter(order__in=cls._delivered_orders(start, end))
            .values("product_name", "pack")
            .annotate(units=Sum("qty"), revenue=Sum(F("rate") * F("qty")))
            .order_by("-revenue")
        )

    @classmethod
    def render_xlsx(cls, data: dict) -> bytes:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        ink = "0A1F3C"
        header_fill = PatternFill("solid", fgColor=ink)
        header_font = Font(bold=True, color="FFFFFF")
        total_fill = PatternFill("solid", fgColor="F3EEE2")
        bold = Font(bold=True)
        thin = Side(style="thin", color="DDD5C4")
        money = '"₹"#,##,##0'

        start = date.fromisoformat(data["start"])
        end = date.fromisoformat(data["end"])
        period = data["period"]
        seller = StoreSettings.load()

        def style_header(ws, row, ncols):
            for c in range(1, ncols + 1):
                cell = ws.cell(row=row, column=c)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            ws.row_dimensions[row].height = 22

        def write_table(ws, first_row, headers, rows, money_cols, total_label="Total"):
            for i, h in enumerate(headers, start=1):
                ws.cell(row=first_row, column=i, value=h)
            style_header(ws, first_row, len(headers))
            r = first_row
            for values in rows:
                r += 1
                for i, v in enumerate(values, start=1):
                    cell = ws.cell(row=r, column=i, value=v)
                    cell.border = Border(bottom=thin)
                    if i in money_cols:
                        cell.number_format = money
            if rows:
                total_row = r + 1
                ws.cell(row=total_row, column=1, value=total_label)
                for i in range(2, len(headers) + 1):
                    first_value = rows[0][i - 1]
                    if isinstance(first_value, (int, float)):
                        col = get_column_letter(i)
                        cell = ws.cell(row=total_row, column=i, value=f"=SUM({col}{first_row + 1}:{col}{r})")
                        if i in money_cols:
                            cell.number_format = money
                for i in range(1, len(headers) + 1):
                    ws.cell(row=total_row, column=i).font = bold
                    ws.cell(row=total_row, column=i).fill = total_fill
            ws.freeze_panes = ws.cell(row=first_row + 1, column=1)

        def autosize(ws, min_width=10, max_width=40):
            for col_cells in ws.columns:
                length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
                ws.column_dimensions[get_column_letter(col_cells[0].column)].width = max(
                    min_width, min(max_width, length + 3)
                )

        wb = Workbook()

        # ── Summary ──
        ws = wb.active
        ws.title = "Summary"
        ws["A1"] = f"{seller.resolved_business_name} — Delivered sales ({period.title()})"
        ws["A1"].font = Font(bold=True, size=14, color=ink)
        ws["A2"] = f"{start.strftime('%d %b %Y')} to {end.strftime('%d %b %Y')} · delivered orders only"
        ws["A2"].font = Font(italic=True, color="75777C")

        t = data["totals"]
        kpis = [
            ("Delivered orders", t["orders"], False),
            ("Items sold", t["items"], False),
            ("Total revenue", t["revenue"], True),
            ("Avg order value", t["avg_order_value"], True),
            ("COD revenue", t["cod_revenue"], True),
            ("Online revenue", t["online_revenue"], True),
            ("GST collected", t["cgst"] + t["sgst"], True),
        ]
        for i, (label, value, is_money) in enumerate(kpis, start=4):
            ws.cell(row=i, column=1, value=label).font = bold
            cell = ws.cell(row=i, column=2, value=value)
            if is_money:
                cell.number_format = money

        table_row = 4 + len(kpis) + 1
        headers = [
            "Period", "Orders", "Items", "Subtotal", "Discount", "CGST", "SGST",
            "Delivery fee", "Revenue", "COD", "Online",
        ]
        rows = [
            [
                r["label"], r["orders"], r["items"], r["subtotal"], r["discount"], r["cgst"], r["sgst"],
                r["delivery_fee"], r["revenue"], r["cod_revenue"], r["online_revenue"],
            ]
            for r in data["rows"]
        ]
        write_table(ws, table_row, headers, rows, money_cols={4, 5, 6, 7, 8, 9, 10, 11})
        autosize(ws)
        ws.column_dimensions["A"].width = 28

        # ── Orders ──
        ws = wb.create_sheet("Delivered orders")
        headers = [
            "Order #", "Delivered on", "Customer", "Mobile", "Payment", "Items", "Subtotal",
            "Discount", "CGST", "SGST", "Delivery fee", "Total",
        ]
        rows = [
            [
                o["order_number"], o["delivered_at"].replace(tzinfo=None), o["customer"], o["mobile"],
                o["payment_mode"], o["items"], o["subtotal"], o["discount"], o["cgst"], o["sgst"],
                o["delivery_fee"], o["total"],
            ]
            for o in data["_orders"]
        ]
        write_table(ws, 1, headers, rows, money_cols={7, 8, 9, 10, 11, 12})
        for r in range(2, len(rows) + 2):
            ws.cell(row=r, column=2).number_format = "dd-mmm-yyyy hh:mm"
        autosize(ws)
        ws.column_dimensions["B"].width = 20

        # ── Products ──
        ws = wb.create_sheet("Products sold")
        rows = [
            [p["product_name"], p["pack"], p["units"], p["revenue"]]
            for p in cls.top_products(start, end)
        ]
        write_table(ws, 1, ["Product", "Pack", "Units sold", "Revenue"], rows, money_cols={4})
        autosize(ws)

        buffer = BytesIO()
        wb.save(buffer)
        return buffer.getvalue()
